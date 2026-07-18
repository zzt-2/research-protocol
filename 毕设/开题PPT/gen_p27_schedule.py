#!/usr/bin/env python3
"""P27: 预期成果与进度安排 — 成果卡片 + 甘特图"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import os

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
FONT = "微软雅黑"

C_DARK   = RGBColor(0x33, 0x33, 0x33)
C_MED    = RGBColor(0x55, 0x55, 0x55)
C_WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
C_BLUE   = RGBColor(0x2E, 0x75, 0xB6)
C_BLUE_L = RGBColor(0xDE, 0xEB, 0xF7)
C_BLUE_B = RGBColor(0x9D, 0xC3, 0xE6)
C_ORG    = RGBColor(0xC5, 0x5A, 0x11)
C_ORG_L  = RGBColor(0xFF, 0xF8, 0xE1)
C_ORG_B  = RGBColor(0xFB, 0xE5, 0xD6)
C_GRAY_L = RGBColor(0xF5, 0xF5, 0xF5)
C_GRAY_B = RGBColor(0xDD, 0xDD, 0xDD)

OUT = "/mnt/d/code/study/research-protocol/毕设/开题PPT"
RECT = 1

# Gantt chart image path
GANTT = "/mnt/d/code/study/research-protocol/毕设/开题PPT/gantt.png"


def add_run(p, text, sz=18, bold=False, color=C_MED):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(sz)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = FONT


def main():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    box_l = Inches(0.5)
    box_w = Inches(12.3)

    # ── Title ──
    tb = slide.shapes.add_textbox(box_l, Inches(0.2), box_w, Inches(0.5))
    tb.text_frame.word_wrap = True
    p = tb.text_frame.paragraphs[0]
    add_run(p, "4.1  预期成果与进度安排", sz=22, bold=True, color=C_DARK)

    # ── Section label: 预期成果 ──
    sec = slide.shapes.add_textbox(box_l, Inches(0.85), box_w, Inches(0.4))
    sec.text_frame.word_wrap = True
    p = sec.text_frame.paragraphs[0]
    add_run(p, "预期成果", sz=18, bold=True, color=C_BLUE)

    # ── 3 outcome cards ──
    outcomes = [
        ("拟建立 GG 湍流下信道估计误差对接收性能影响的量化分析框架，给出面向载波同步的精度设计方法",
         C_BLUE, C_BLUE_L, C_BLUE_B),
        ("拟建立湍流下载波同步方法的性能分析框架，给出分湍流条件的方法选择与参数设计方法",
         C_ORG, C_ORG_L, C_ORG_B),
        ("从跨模块联合处理、编码辅助恢复等方向探索湍流下接收端信号处理方法",
         C_BLUE, C_BLUE_L, C_BLUE_B),
    ]

    cw = Inches(3.9)
    gap = Inches(0.25)
    x1 = Inches(0.5)
    x2 = x1 + cw + gap
    x3 = x2 + cw + gap
    cy = Inches(1.35)
    ch = Inches(1.8)

    for i, (text, accent, fill, border) in enumerate(outcomes):
        x = [x1, x2, x3][i]
        # Card
        card = slide.shapes.add_shape(RECT, x, cy, cw, ch)
        card.fill.solid()
        card.fill.fore_color.rgb = fill
        card.line.color.rgb = border
        card.line.width = Pt(1)
        # Left accent
        s = slide.shapes.add_shape(RECT, x, cy, Inches(0.06), ch)
        s.fill.solid()
        s.fill.fore_color.rgb = accent
        s.line.fill.background()
        # Text
        tb = slide.shapes.add_textbox(x + Inches(0.2), cy + Inches(0.15),
                                       cw - Inches(0.35), ch - Inches(0.3))
        tb.text_frame.word_wrap = True
        p = tb.text_frame.paragraphs[0]
        add_run(p, text, sz=18, color=C_DARK)

    # ── Section label: 进度安排 ──
    sec2 = slide.shapes.add_textbox(box_l, Inches(3.4), box_w, Inches(0.4))
    sec2.text_frame.word_wrap = True
    p = sec2.text_frame.paragraphs[0]
    add_run(p, "进度安排", sz=18, bold=True, color=C_BLUE)

    # ── Gantt chart ──
    if os.path.exists(GANTT):
        slide.shapes.add_picture(GANTT, Inches(0), Inches(3.9), Inches(13.333))
    else:
        # Placeholder if image not found
        ph = slide.shapes.add_shape(RECT, Inches(0.5), Inches(3.9), box_w, Inches(3.3))
        ph.fill.solid()
        ph.fill.fore_color.rgb = C_GRAY_L
        ph.line.color.rgb = C_GRAY_B
        ph.line.width = Pt(1)
        tb = slide.shapes.add_textbox(Inches(5), Inches(5.2), Inches(3), Inches(0.5))
        tb.text_frame.word_wrap = True
        p = tb.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        add_run(p, "[甘特图]", sz=18, color=C_MED)

    out = os.path.join(OUT, "P27-预期成果与进度安排v1.pptx")
    prs.save(out)
    print(f"Saved → {out}")


if __name__ == "__main__":
    main()
