#!/usr/bin/env python3
"""P16(Ch3) + P18(Ch4): 18pt最小字号, 微软雅黑, 色框+箭头"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import os

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
FONT = "微软雅黑"

C_DARK   = RGBColor(0x33, 0x33, 0x33)
C_MED    = RGBColor(0x44, 0x44, 0x44)
C_BLUE   = RGBColor(0x2E, 0x75, 0xB6)
C_BLUE_L = RGBColor(0xDE, 0xEB, 0xF7)
C_BLUE_B = RGBColor(0x9D, 0xC3, 0xE6)
C_BLUE_D = RGBColor(0x1F, 0x4E, 0x79)
C_ORG    = RGBColor(0xC5, 0x5A, 0x11)
C_ORG_L  = RGBColor(0xFF, 0xF8, 0xE1)
C_ORG_B  = RGBColor(0xFB, 0xE5, 0xD6)
C_ORG_D  = RGBColor(0x84, 0x3C, 0x0B)
C_GRAY_L = RGBColor(0xF5, 0xF5, 0xF5)
C_GRAY_B = RGBColor(0xDD, 0xDD, 0xDD)

OUT = "/mnt/d/code/study/research-protocol/毕设/开题PPT"


def run(p, text, sz=18, bold=False, color=C_MED):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(sz)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = FONT


def add_title(slide, text):
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(12.3), Inches(0.7))
    tb.text_frame.word_wrap = True
    p = tb.text_frame.paragraphs[0]
    run(p, text, sz=22, bold=True, color=C_DARK)


def add_box(slide, top, h, label, bullets, fill, border, label_c, bullet_c):
    """Rounded box with label + bullet lines. Returns bottom y."""
    shape = slide.shapes.add_shape(5, Inches(0.5), top, Inches(12.3), h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = border
    shape.line.width = Pt(1.5)

    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.25)
    tf.margin_right = Inches(0.25)
    tf.margin_top = Inches(0.12)
    tf.margin_bottom = Inches(0.08)

    p = tf.paragraphs[0]
    run(p, label, sz=18, bold=True, color=label_c)
    p.space_after = Pt(6)

    for b in bullets:
        pb = tf.add_paragraph()
        run(pb, f"  {b}", sz=18, color=bullet_c)
        pb.space_after = Pt(3)

    return top + h


def add_arrow(slide, top, color):
    shape = slide.shapes.add_shape(13, Inches(6.47), top, Inches(0.4), Inches(0.3))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return top + Inches(0.3)


def add_highlight(slide, top, text, fill, border, text_c):
    shape = slide.shapes.add_shape(5, Inches(0.5), top, Inches(12.3), Inches(0.65))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = border
    shape.line.width = Pt(2)

    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.25)
    tf.margin_top = Inches(0.1)
    tf.margin_bottom = Inches(0.05)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run(p, text, sz=18, bold=True, color=text_c)


def main():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    # ═══════════════════════════════════════
    # P16: Ch3
    # ═══════════════════════════════════════
    s1 = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s1,
        "信道估计误差对 QPSK 影响可忽略，对 16-QAM 高度敏感")

    y = add_box(s1, Inches(1.0), Inches(1.15),
        "针对问题",
        ["信道估计误差如何级联影响载波同步性能，缺乏量化分析",
         "不同估计方法精度差异及其对调制格式的影响尚不明确"],
        C_GRAY_L, C_GRAY_B, C_DARK, C_MED)

    y = add_arrow(s1, y + Inches(0.08), C_BLUE)

    y = add_box(s1, y + Inches(0.08), Inches(1.55),
        "核心发现",
        ["QPSK 符号判决对信道缩放不敏感，工作 SNR 下估计误差影响可忽略",
         "16-QAM 判决依赖信号幅度，估计误差引入缩放偏移，BER 退化显著",
         "强湍流下 E[1/h] 发散，前馈方法的统计最优性前提不成立"],
        C_BLUE_L, C_BLUE_B, C_BLUE, C_BLUE_D)

    y = add_arrow(s1, y + Inches(0.08), C_BLUE)

    y = add_box(s1, y + Inches(0.08), Inches(1.0),
        "产出 (IP1)",
        ["建立估计误差对接收性能影响的量化分析框架",
         "给出面向同步的精度设计准则，分湍流条件和调制格式给出精度需求"],
        C_BLUE_L, C_BLUE_B, C_BLUE, C_BLUE)

    add_arrow(s1, y + Inches(0.08), C_BLUE)

    add_highlight(s1, y + Inches(0.4),
        "调制格式决定精度需求——QPSK 鲁棒而 16-QAM 敏感，精度准则需区别对待",
        C_BLUE_L, C_BLUE, C_BLUE_D)

    # ═══════════════════════════════════════
    # P18: Ch4
    # ═══════════════════════════════════════
    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s2,
        "不同载波同步方法对湍流参数灵敏度差异显著，分条件设计准则指导工程配置")

    y = add_box(s2, Inches(1.0), Inches(1.15),
        "针对问题",
        ["湍流下载波同步方法缺乏系统性性能分析与设计指导",
         "经典设计公式基于白噪声假设，湍流适配准则不完善"],
        C_GRAY_L, C_GRAY_B, C_DARK, C_MED)

    y = add_arrow(s2, y + Inches(0.08), C_ORG)

    y = add_box(s2, y + Inches(0.08), Inches(1.55),
        "方法对比与失效机制",
        ["VV/BPS 前馈方法在弱、中湍流下可用，强湍流深衰落导致相位解卷绕失败",
         "DPLL 在合理参数配置下全湍流范围稳定跟踪，强湍流下性能最优",
         "参数灵敏度差异是设计准则的核心依据"],
        C_ORG_L, C_ORG_B, C_ORG, C_ORG_D)

    y = add_arrow(s2, y + Inches(0.08), C_ORG)

    y = add_box(s2, y + Inches(0.08), Inches(1.0),
        "产出 (IP2)",
        ["分湍流条件的方法选择与参数设计准则",
         "三层次递进：性能分析 → 参数设计 → 联合处理"],
        C_ORG_L, C_ORG_B, C_ORG, C_ORG)

    add_arrow(s2, y + Inches(0.08), C_ORG)

    add_highlight(s2, y + Inches(0.4),
        "设计准则核心价值在于分条件匹配方法与参数",
        C_ORG_L, C_ORG, C_ORG_D)

    out = os.path.join(OUT, "P16-P18-Ch3-Ch4-概述v5.pptx")
    prs.save(out)
    print(f"Saved {len(prs.slides)} slides → {out}")


if __name__ == "__main__":
    main()
