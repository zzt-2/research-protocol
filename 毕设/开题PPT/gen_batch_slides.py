#!/usr/bin/env python3
"""批量生成开题PPT页面: P12(概览), P14(Ch2), P16(Ch3), P18(Ch4), P20(Ch5), P28(预期成果)"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import os

# ── Constants ──
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
FONT = "微软雅黑"

# Colors
C_DARK_GRAY = RGBColor(0x44, 0x44, 0x44)
C_MED_GRAY  = RGBColor(0x66, 0x66, 0x66)
C_LIGHT_GRAY = RGBColor(0xF2, 0xF2, 0xF2)
C_BORDER     = RGBColor(0xCC, 0xCC, 0xCC)
C_BLUE       = RGBColor(0x2E, 0x75, 0xB6)
C_LIGHT_BLUE = RGBColor(0xDE, 0xEB, 0xF7)
C_BLUE_BORDER = RGBColor(0x9D, 0xC3, 0xE6)
C_BLUE_DARK  = RGBColor(0x1F, 0x4E, 0x79)
C_ORANGE     = RGBColor(0xC5, 0x5A, 0x11)
C_LIGHT_ORANGE = RGBColor(0xFF, 0xF8, 0xE1)
C_ORANGE_BORDER = RGBColor(0xFB, 0xE5, 0xD6)
C_ORANGE_DARK = RGBColor(0x84, 0x3C, 0x0B)

IMG_DIR = "/mnt/d/code/study/research-protocol/毕设/开题报告/figures/png"
OUT_DIR = "/mnt/d/code/study/research-protocol/毕设/开题PPT"


def add_title(slide, text, size=22):
    txBox = slide.shapes.add_textbox(Inches(0.6), Inches(0.3), Inches(12.1), Inches(0.8))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = True
    run.font.color.rgb = C_DARK_GRAY
    run.font.name = FONT


def add_paragraph(slide, left, top, width, height, text, size=12, color=C_MED_GRAY):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = FONT
    return txBox


def add_block(slide, left, top, w, h, fill, border, title, body_lines,
              title_color=C_DARK_GRAY, body_color=C_MED_GRAY,
              title_size=13, body_size=11):
    shape = slide.shapes.add_shape(1, left, top, w, h)  # RECTANGLE
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = border
    shape.line.width = Pt(1)

    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.15)
    tf.margin_right = Inches(0.15)
    tf.margin_top = Inches(0.12)
    tf.margin_bottom = Inches(0.1)

    # Title
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    run.font.size = Pt(title_size)
    run.font.bold = True
    run.font.color.rgb = title_color
    run.font.name = FONT
    p.space_after = Pt(6)

    # Body lines
    for line in body_lines:
        p2 = tf.add_paragraph()
        run2 = p2.add_run()
        run2.text = line
        run2.font.size = Pt(body_size)
        run2.font.color.rgb = body_color
        run2.font.name = FONT
        p2.space_after = Pt(2)
        p2.line_spacing = Pt(body_size + 5)

    return shape


def add_image(slide, path, left, top, width):
    if os.path.exists(path):
        slide.shapes.add_picture(path, left, top, width)
    else:
        print(f"  WARNING: {path} not found")


def make_content_slide(prs, title, overview, img_filename):
    """Standard content slide: title + overview text + large drawio image"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, title)
    add_paragraph(slide, Inches(0.6), Inches(1.2), Inches(12.1), Inches(0.9),
                  overview, size=12, color=C_MED_GRAY)
    img_path = os.path.join(IMG_DIR, img_filename)
    # Center image, let height auto-scale
    add_image(slide, img_path, Inches(1.0), Inches(2.3), Inches(11.3))
    return slide


def main():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    # ═══════════════════════════════════════
    # P12: 研究内容概览 (2×2 blocks)
    # ═══════════════════════════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "研究内容围绕湍流信道建模、信道估计、信号处理方法与硬件验证四个问题展开")

    bw, bh = Inches(5.85), Inches(2.5)
    x1, x2 = Inches(0.5), Inches(6.85)
    y1, y2 = Inches(1.3), Inches(4.1)

    # Ch2 - neutral
    add_block(slide, x1, y1, bw, bh, C_LIGHT_GRAY, C_BORDER,
              "Ch2  星地激光通信系统与信道模型",
              ["针对星地链路的系统建模与湍流参数量化需求，建立 QPSK 相干检测系统模型与",
               "Gamma-Gamma 湍流信道模型，完成链路预算分析，为后续章节提供统一参数基础。"],
              C_DARK_GRAY, C_MED_GRAY)

    # Ch3 - blue
    add_block(slide, x2, y1, bw, bh, C_LIGHT_BLUE, C_BLUE_BORDER,
              "Ch3  大气湍流信道估计与接收性能分析",
              ["针对信道估计误差对接收性能的量化影响尚不明确的问题，对比分析信道估计方法",
               "精度，研究误差对接收性能的影响机制，给出方法可用性边界与面向同步的精度设计准则。"],
              C_BLUE, C_BLUE_DARK)

    # Ch4 - orange
    add_block(slide, x1, y2, bw, bh, C_LIGHT_ORANGE, C_ORANGE_BORDER,
              "Ch4  湍流下相干接收端信号处理方法研究",
              ["针对湍流下载波同步方法缺乏系统性分析与设计指导的问题，分析性能边界与失效",
               "机制，建立方法选择与参数设计准则，并从信号处理方法层面开展探索。"],
              C_ORANGE, C_ORANGE_DARK)

    # Ch5 - neutral
    add_block(slide, x2, y2, bw, bh, C_LIGHT_GRAY, C_BORDER,
              "Ch5  接收端信号处理链关键技术实现与验证",
              ["针对信号处理方法的工程可实现性与综合性能评估问题，开展接收端信号处理链",
               "关键技术的硬件实现与验证。"],
              C_DARK_GRAY, C_MED_GRAY)

    # ═══════════════════════════════════════
    # P14: Ch2
    # ═══════════════════════════════════════
    make_content_slide(prs,
        "建立 QPSK 相干检测系统模型与 Gamma-Gamma 湍流信道模型，为后续章节提供统一参数基础",
        "针对星地链路的系统建模与湍流参数量化需求，建立 QPSK 相干检测系统模型与 Gamma-Gamma 湍流信道模型，完成链路预算分析，为后续章节提供统一参数基础。",
        "P15-Ch2-链路预算.png")

    # ═══════════════════════════════════════
    # P16: Ch3
    # ═══════════════════════════════════════
    make_content_slide(prs,
        "对比分析估计方法精度，给出方法可用性边界与面向同步的精度设计准则",
        "针对信道估计误差对接收性能的量化影响尚不明确的问题，对比分析信道估计方法精度，研究误差对接收性能的影响机制，给出方法可用性边界与面向同步的精度设计准则。",
        "P17-Ch3-精度准则.png")

    # ═══════════════════════════════════════
    # P18: Ch4
    # ═══════════════════════════════════════
    make_content_slide(prs,
        "分析载波同步方法性能边界，建立方法选择与参数设计准则，并探索信号处理方法",
        "针对湍流下载波同步方法缺乏系统性分析与设计指导的问题，分析性能边界与失效机制，建立方法选择与参数设计准则，并从信号处理方法层面开展探索。",
        "P19-Ch4-准则与探索.png")

    # ═══════════════════════════════════════
    # P20: Ch5
    # ═══════════════════════════════════════
    make_content_slide(prs,
        "开展接收端信号处理链关键技术的硬件实现与验证",
        "针对信号处理方法的工程可实现性与综合性能评估问题，开展接收端信号处理链关键技术的硬件实现与验证。",
        "P20-Ch5.png")

    # ═══════════════════════════════════════
    # P28: 预期成果与进度 (3 tier blocks + Gantt)
    # ═══════════════════════════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "预期成果涵盖理论分析框架、方法探索与学术产出三个层面")

    tw, th = Inches(3.9), Inches(2.1)
    ty = Inches(1.15)

    # Tier 1: 理论 (blue)
    add_block(slide, Inches(0.5), ty, tw, th, C_LIGHT_BLUE, C_BLUE_BORDER,
              "理论层面",
              ["① 拟建立 GG 湍流下信道估计误差对接收性能影响的量化分析框架",
               "② 拟建立湍流下载波同步方法的性能分析框架"],
              C_BLUE, C_BLUE_DARK, body_size=11)

    # Tier 2: 方法探索 (orange)
    add_block(slide, Inches(4.6), ty, tw, th, C_LIGHT_ORANGE, C_ORANGE_BORDER,
              "方法探索层面",
              ["③ 拟从跨模块联合处理、编码辅助恢复等方向",
               "   探索湍流下接收端信号处理方法"],
              C_ORANGE, C_ORANGE_DARK, body_size=11)

    # Tier 3: 学术产出 (neutral)
    add_block(slide, Inches(8.7), ty, tw, th, C_LIGHT_GRAY, C_BORDER,
              "学术产出",
              ["④ 发表一篇综述论文",
               "⑤ 发表一篇会议论文"],
              C_DARK_GRAY, C_MED_GRAY, body_size=11)

    # Gantt chart
    add_image(slide, os.path.join(IMG_DIR, "P29-Gantt.png"),
              Inches(1.5), Inches(3.6), Inches(10.3))

    # ── Save ──
    out = os.path.join(OUT_DIR, "开题答辩_批量生成_P12-P28.pptx")
    prs.save(out)
    print(f"Saved {len(prs.slides)} slides → {out}")


if __name__ == "__main__":
    main()
