#!/usr/bin/env python3
"""P6: 方向分析 — 信道估计与载波同步方法概览 + 缺口"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import os

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
FONT = "微软雅黑"

C_DARK   = RGBColor(0x33, 0x33, 0x33)
C_MED    = RGBColor(0x55, 0x55, 0x55)
C_WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
C_BLUE   = RGBColor(0x2E, 0x75, 0xB6)
C_BLUE_L = RGBColor(0xDE, 0xEB, 0xF7)
C_BLUE_B = RGBColor(0x9D, 0xC3, 0xE6)
C_BLUE_D = RGBColor(0x1F, 0x4E, 0x79)
C_ORG    = RGBColor(0xC5, 0x5A, 0x11)
C_ORG_L  = RGBColor(0xFF, 0xF8, 0xE1)
C_ORG_B  = RGBColor(0xFB, 0xE5, 0xD6)
C_ORG_D  = RGBColor(0x84, 0x3C, 0x0B)
C_GRAY_L = RGBColor(0xF5, 0xF5, 0xF5)
C_GRAY_B = RGBColor(0xDD, 0xDD, 0xDD)

OUT = "/mnt/d/code/study/research-protocol/毕设/开题PPT"
RECT = 1


def add_run(p, text, sz=18, bold=False, color=C_MED):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(sz)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = FONT


def main():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    box_l = Inches(0.5)
    box_w = Inches(12.3)

    # ── Title ──
    tb = slide.shapes.add_textbox(box_l, Inches(0.2), box_w, Inches(0.6))
    tb.text_frame.word_wrap = True
    p = tb.text_frame.paragraphs[0]
    add_run(p, "信道估计与载波同步在 FSO 湍流中均存在覆盖缺口，二者构成上下游约束",
            sz=22, bold=True, color=C_DARK)

    # ── Column dims ──
    col_w = Inches(5.9)
    col_gap = Inches(0.5)
    col_l_x = box_l
    col_r_x = box_l + col_w + col_gap

    hdr_y = Inches(0.95)
    hdr_h = Inches(0.55)

    # ── LEFT: 信道估计 (blue) ──
    lh = slide.shapes.add_shape(RECT, col_l_x, hdr_y, col_w, hdr_h)
    lh.fill.solid()
    lh.fill.fore_color.rgb = C_BLUE_D
    lh.line.fill.background()

    tb = slide.shapes.add_textbox(col_l_x + Inches(0.15), hdr_y + Inches(0.1),
                                   col_w - Inches(0.3), hdr_h - Inches(0.2))
    tb.text_frame.word_wrap = True
    p = tb.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    add_run(p, "信道估计", sz=20, bold=True, color=C_WHITE)

    # Left method cards
    methods_l = [
        ("最小二乘法", "结构简洁"),
        ("最小均方误差", "先验辅助降噪"),
        ("卡尔曼滤波", "递推跟踪时变"),
        ("深度学习", "缓解导频依赖"),
    ]

    mc_w = col_w - Inches(0.2)
    mc_h = Inches(0.6)
    mc_gap = Inches(0.1)
    mc_x = col_l_x + Inches(0.1)
    mc_y0 = hdr_y + hdr_h + Inches(0.15)

    for i, (name, desc) in enumerate(methods_l):
        y = mc_y0 + i * (mc_h + mc_gap)
        c = slide.shapes.add_shape(RECT, mc_x, y, mc_w, mc_h)
        c.fill.solid()
        c.fill.fore_color.rgb = C_WHITE
        c.line.color.rgb = C_BLUE_B
        c.line.width = Pt(1)
        s = slide.shapes.add_shape(RECT, mc_x, y, Inches(0.06), mc_h)
        s.fill.solid()
        s.fill.fore_color.rgb = C_BLUE
        s.line.fill.background()
        tb = slide.shapes.add_textbox(mc_x + Inches(0.2), y + Inches(0.12),
                                       mc_w - Inches(0.35), mc_h - Inches(0.24))
        tb.text_frame.word_wrap = True
        p = tb.text_frame.paragraphs[0]
        add_run(p, name, sz=18, bold=True, color=C_BLUE)
        add_run(p, f"  ·  {desc}", sz=18, color=C_MED)

    # Left gap box
    gap_y = mc_y0 + 4 * (mc_h + mc_gap) + Inches(0.15)
    gap_h = Inches(1.2)

    gb = slide.shapes.add_shape(RECT, mc_x, gap_y, mc_w, gap_h)
    gb.fill.solid()
    gb.fill.fore_color.rgb = C_BLUE_L
    gb.line.color.rgb = C_BLUE
    gb.line.width = Pt(1.5)
    gs = slide.shapes.add_shape(RECT, mc_x, gap_y, Inches(0.08), gap_h)
    gs.fill.solid()
    gs.fill.fore_color.rgb = C_BLUE
    gs.line.fill.background()

    tb = slide.shapes.add_textbox(mc_x + Inches(0.25), gap_y + Inches(0.15),
                                   mc_w - Inches(0.4), gap_h - Inches(0.3))
    tb.text_frame.word_wrap = True
    p = tb.text_frame.paragraphs[0]
    add_run(p, "缺口：", sz=18, bold=True, color=C_BLUE)
    p2 = tb.text_frame.add_paragraph()
    add_run(p2, "估计精度与下游载波同步需求缺乏定量关联", sz=18, color=C_DARK)

    # ── RIGHT: 载波同步 (orange) ──
    rh = slide.shapes.add_shape(RECT, col_r_x, hdr_y, col_w, hdr_h)
    rh.fill.solid()
    rh.fill.fore_color.rgb = C_ORG_D
    rh.line.fill.background()

    tb = slide.shapes.add_textbox(col_r_x + Inches(0.15), hdr_y + Inches(0.1),
                                   col_w - Inches(0.3), hdr_h - Inches(0.2))
    tb.text_frame.word_wrap = True
    p = tb.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    add_run(p, "载波同步", sz=20, bold=True, color=C_WHITE)

    methods_r = [
        ("VV 算法", "QPSK 主流前馈方法"),
        ("盲相位搜索", "高阶调制适应性好"),
        ("数字锁相环", "闭环跟踪相位"),
        ("深度学习", "湍流适用性空白"),
    ]

    rc_x = col_r_x + Inches(0.1)

    for i, (name, desc) in enumerate(methods_r):
        y = mc_y0 + i * (mc_h + mc_gap)
        c = slide.shapes.add_shape(RECT, rc_x, y, mc_w, mc_h)
        c.fill.solid()
        c.fill.fore_color.rgb = C_WHITE
        c.line.color.rgb = C_ORG_B
        c.line.width = Pt(1)
        s = slide.shapes.add_shape(RECT, rc_x, y, Inches(0.06), mc_h)
        s.fill.solid()
        s.fill.fore_color.rgb = C_ORG
        s.line.fill.background()
        tb = slide.shapes.add_textbox(rc_x + Inches(0.2), y + Inches(0.12),
                                       mc_w - Inches(0.35), mc_h - Inches(0.24))
        tb.text_frame.word_wrap = True
        p = tb.text_frame.paragraphs[0]
        add_run(p, name, sz=18, bold=True, color=C_ORG)
        add_run(p, f"  ·  {desc}", sz=18, color=C_MED)

    # Right gap box
    rgb = slide.shapes.add_shape(RECT, rc_x, gap_y, mc_w, gap_h)
    rgb.fill.solid()
    rgb.fill.fore_color.rgb = C_ORG_L
    rgb.line.color.rgb = C_ORG
    rgb.line.width = Pt(1.5)
    rs = slide.shapes.add_shape(RECT, rc_x, gap_y, Inches(0.08), gap_h)
    rs.fill.solid()
    rs.fill.fore_color.rgb = C_ORG
    rs.line.fill.background()

    tb = slide.shapes.add_textbox(rc_x + Inches(0.25), gap_y + Inches(0.15),
                                   mc_w - Inches(0.4), gap_h - Inches(0.3))
    tb.text_frame.word_wrap = True
    p = tb.text_frame.paragraphs[0]
    add_run(p, "缺口：", sz=18, bold=True, color=C_ORG)
    p2 = tb.text_frame.add_paragraph()
    add_run(p2, "经典设计公式基于白噪声，湍流适配参数设计准则不完善", sz=18, color=C_DARK)

    # ── Bottom bridge ──
    bridge_y = gap_y + gap_h + Inches(0.2)
    bridge_h = Inches(0.65)

    bridge = slide.shapes.add_shape(RECT, box_l, bridge_y, box_w, bridge_h)
    bridge.fill.solid()
    bridge.fill.fore_color.rgb = C_GRAY_L
    bridge.line.color.rgb = C_BLUE_B
    bridge.line.width = Pt(1.5)

    bs1 = slide.shapes.add_shape(RECT, box_l, bridge_y, Inches(0.08), bridge_h)
    bs1.fill.solid()
    bs1.fill.fore_color.rgb = C_BLUE
    bs1.line.fill.background()

    bs2 = slide.shapes.add_shape(RECT, box_l + box_w - Inches(0.08), bridge_y,
                                  Inches(0.08), bridge_h)
    bs2.fill.solid()
    bs2.fill.fore_color.rgb = C_ORG
    bs2.line.fill.background()

    tb = slide.shapes.add_textbox(box_l + Inches(0.3), bridge_y + Inches(0.1),
                                   box_w - Inches(0.6), bridge_h - Inches(0.2))
    tb.text_frame.word_wrap = True
    p = tb.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    add_run(p, "信道估计精度 ", sz=18, bold=True, color=C_BLUE)
    add_run(p, "→ ", sz=20, bold=True, color=C_DARK)
    add_run(p, " 影响载波同步起点", sz=18, bold=True, color=C_ORG)

    out = os.path.join(OUT, "P06-方向分析v1.pptx")
    prs.save(out)
    print(f"Saved → {out}")


if __name__ == "__main__":
    main()
