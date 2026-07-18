#!/usr/bin/env python3
"""P23: 创新点 + 右侧可扩展方向面板"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
import os

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
FONT = "微软雅黑"

C_DARK   = RGBColor(0x33, 0x33, 0x33)
C_MED    = RGBColor(0x55, 0x55, 0x55)
C_WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
C_BLUE   = RGBColor(0x2E, 0x75, 0xB6)
C_BLUE_L = RGBColor(0xDE, 0xEB, 0xF7)
C_BLUE_B = RGBColor(0x5B, 0x9B, 0xD5)
C_ORG    = RGBColor(0xC5, 0x5A, 0x11)
C_ORG_L  = RGBColor(0xFF, 0xF8, 0xE1)
C_ORG_B  = RGBColor(0xED, 0x7D, 0x31)
C_GRAY_L = RGBColor(0xF5, 0xF5, 0xF5)
C_GRAY_B = RGBColor(0xCC, 0xCC, 0xCC)
C_GRAY_T = RGBColor(0x99, 0x99, 0x99)

OUT = "/mnt/d/code/study/research-protocol/毕设/开题PPT"
RECT = 1


def add_run(p, text, sz=18, bold=False, color=C_MED):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(sz)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = FONT


def set_dash(shape):
    ln = shape.line._ln
    ln.append(ln.makeelement(qn('a:prstDash'), {'val': 'dash'}))


def main():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # ── Layout ──
    card_l = Inches(0.5)
    card_w = Inches(8.3)
    panel_l = Inches(9.3)
    panel_w = Inches(3.5)

    # ── Title ──
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(12.3), Inches(0.5))
    tb.text_frame.word_wrap = True
    p = tb.text_frame.paragraphs[0]
    add_run(p, "3.5  创新点", sz=22, bold=True, color=C_DARK)

    # ── Intro ──
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.85), Inches(12.3), Inches(0.4))
    tb.text_frame.word_wrap = True
    p = tb.text_frame.paragraphs[0]
    add_run(p, "本课题拟围绕两个创新点开展研究", sz=18, bold=True, color=C_MED)

    # ── Card 1: IP1 (Blue) ──
    c1_y = Inches(1.45)
    c1_h = Inches(2.1)

    card1 = slide.shapes.add_shape(RECT, card_l, c1_y, card_w, c1_h)
    card1.fill.solid()
    card1.fill.fore_color.rgb = C_BLUE_L
    card1.line.color.rgb = C_BLUE_B
    card1.line.width = Pt(1.5)

    bar1 = slide.shapes.add_shape(RECT, card_l, c1_y, Inches(0.1), c1_h)
    bar1.fill.solid()
    bar1.fill.fore_color.rgb = C_BLUE
    bar1.line.fill.background()

    tb = slide.shapes.add_textbox(card_l + Inches(0.25), c1_y + Inches(0.15),
                                  card_w - Inches(0.4), c1_h - Inches(0.3))
    tb.text_frame.word_wrap = True
    p = tb.text_frame.paragraphs[0]
    add_run(p, "创新点(1)  ", sz=14, bold=True, color=C_BLUE)
    add_run(p, "拟建立", sz=14, bold=True, color=C_MED)
    p2 = tb.text_frame.add_paragraph()
    add_run(p2, "湍流FSO信道估计误差对接收性能影响的分析模型", sz=18, color=C_DARK)
    p3 = tb.text_frame.add_paragraph()
    add_run(p3, "· 推导误差传播边界与前馈方法可用性判据", sz=18, color=C_MED)
    p4 = tb.text_frame.add_paragraph()
    add_run(p4, "· 建立面向不同调制格式和湍流条件的精度设计方法", sz=18, color=C_MED)

    # ── Connector label ──
    conn_y = c1_y + c1_h + Inches(0.05)
    tb = slide.shapes.add_textbox(card_l + Inches(2), conn_y,
                                  Inches(4.5), Inches(0.3))
    tb.text_frame.word_wrap = True
    p = tb.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    add_run(p, "精度需求为同步设计提供定量约束", sz=13, color=C_GRAY_T)

    # ── Card 2: IP2 (Orange) ──
    c2_y = conn_y + Inches(0.35)
    c2_h = Inches(2.5)

    card2 = slide.shapes.add_shape(RECT, card_l, c2_y, card_w, c2_h)
    card2.fill.solid()
    card2.fill.fore_color.rgb = C_ORG_L
    card2.line.color.rgb = C_ORG_B
    card2.line.width = Pt(1.5)

    bar2 = slide.shapes.add_shape(RECT, card_l, c2_y, Inches(0.1), c2_h)
    bar2.fill.solid()
    bar2.fill.fore_color.rgb = C_ORG
    bar2.line.fill.background()

    tb = slide.shapes.add_textbox(card_l + Inches(0.25), c2_y + Inches(0.15),
                                  card_w - Inches(0.4), c2_h - Inches(0.3))
    tb.text_frame.word_wrap = True
    p = tb.text_frame.paragraphs[0]
    add_run(p, "创新点(2)  ", sz=14, bold=True, color=C_ORG)
    add_run(p, "拟研究", sz=14, bold=True, color=C_MED)
    p2 = tb.text_frame.add_paragraph()
    add_run(p2, "湍流下载波同步方法的性能分析与参数设计", sz=18, color=C_DARK)
    p3 = tb.text_frame.add_paragraph()
    add_run(p3, "· 建立系统性性能分析模型", sz=18, color=C_MED)
    p4 = tb.text_frame.add_paragraph()
    add_run(p4, "· 给出分湍流条件的方法选择与参数设计方法", sz=18, color=C_MED)
    p5 = tb.text_frame.add_paragraph()
    add_run(p5, "· 从联合处理、编码辅助恢复等方向开展方法探索", sz=18, color=C_MED)

    # ── Right Panel: 可扩展方向 ──
    panel_y = c1_y
    panel_h = c2_y + c2_h - c1_y

    # Panel background (dashed)
    panel = slide.shapes.add_shape(RECT, panel_l, panel_y, panel_w, panel_h)
    panel.fill.solid()
    panel.fill.fore_color.rgb = C_GRAY_L
    panel.line.color.rgb = C_GRAY_B
    panel.line.width = Pt(1.5)
    set_dash(panel)

    # Panel title
    tb = slide.shapes.add_textbox(panel_l + Inches(0.15), panel_y + Inches(0.15),
                                  panel_w - Inches(0.3), Inches(0.35))
    tb.text_frame.word_wrap = True
    p = tb.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    add_run(p, "可扩展方向", sz=18, bold=True, color=C_GRAY_T)

    # Direction items
    directions = [
        "基于湍流预测的前馈补偿",
        "多孔径分集接收",
        "湍流自适应调制",
        "编码辅助载波恢复",
        "数据驱动的信号处理方法",
    ]

    item_w = panel_w - Inches(0.3)
    item_h = Inches(0.65)
    item_gap = Inches(0.12)
    item_x = panel_l + Inches(0.15)
    item_y0 = panel_y + Inches(0.6)

    for i, d in enumerate(directions):
        y = item_y0 + i * (item_h + item_gap)
        ic = slide.shapes.add_shape(RECT, item_x, y, item_w, item_h)
        ic.fill.solid()
        ic.fill.fore_color.rgb = C_WHITE
        ic.line.color.rgb = C_GRAY_B
        ic.line.width = Pt(0.75)
        s = slide.shapes.add_shape(RECT, item_x, y, Inches(0.04), item_h)
        s.fill.solid()
        s.fill.fore_color.rgb = C_GRAY_T
        s.line.fill.background()
        tb = slide.shapes.add_textbox(item_x + Inches(0.15), y + Inches(0.15),
                                      item_w - Inches(0.25), item_h - Inches(0.3))
        tb.text_frame.word_wrap = True
        p = tb.text_frame.paragraphs[0]
        add_run(p, d, sz=18, color=C_MED)

    # ── Arrow between cards and panel ──
    arrow_y = panel_y + panel_h / 2
    tb = slide.shapes.add_textbox(card_l + card_w + Inches(0.02),
                                  arrow_y - Inches(0.15),
                                  Inches(0.46), Inches(0.3))
    tb.text_frame.word_wrap = True
    p = tb.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    add_run(p, "→", sz=20, bold=True, color=C_GRAY_T)

    # "递进" label below arrow
    tb = slide.shapes.add_textbox(card_l + card_w + Inches(0.02),
                                  arrow_y + Inches(0.15),
                                  Inches(0.46), Inches(0.25))
    tb.text_frame.word_wrap = True
    p = tb.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    add_run(p, "递进", sz=13, color=C_GRAY_T)

    # ── Speaker notes ──
    notes_slide = slide.notes_slide
    notes_slide.notes_text_frame.text = (
        "本课题有两个创新点。"
        "第一，建立信道估计误差对接收性能影响的分析模型"
        "——推导误差传播边界和前馈方法的可用性判据，"
        "给出面向不同调制格式和湍流条件的精度设计方法。"
        "第二，研究湍流下载波同步方法的性能分析与参数设计"
        "——建立系统性分析模型，给出分湍流条件的方法选择与参数设计方法，"
        "并从联合处理等方向开展方法探索。"
        "右侧列出了从当前创新点自然延伸的可扩展方向，"
        "包括湍流预测前馈补偿、多孔径分集、自适应调制、"
        "编码辅助载波恢复和数据驱动方法，"
        "这些方向均可在现有框架基础上进一步展开。"
    )

    out = os.path.join(OUT, "P23-创新点+可扩展方向v1.pptx")
    prs.save(out)
    print(f"Saved → {out}")


if __name__ == "__main__":
    main()
