#!/usr/bin/env python3
"""P7: 研究空白分析 — 单框双栏 + 顶部总结 + 底部桥接"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import os

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
FONT = "微软雅黑"

C_DARK   = RGBColor(0x33, 0x33, 0x33)
C_MED    = RGBColor(0x55, 0x55, 0x55)
C_WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
C_BLUE   = RGBColor(0x2E, 0x75, 0xB6)
C_BLUE_L = RGBColor(0xDE, 0xEB, 0xF7)
C_BLUE_B = RGBColor(0x9D, 0xC3, 0xE6)
C_BLUE_D = RGBColor(0x1F, 0x4E, 0x79)
C_ORG    = RGBColor(0xC5, 0x5A, 0x11)
C_ORG_L  = RGBColor(0xFF, 0xF8, 0xE1)
C_ORG_B  = RGBColor(0xFB, 0xE5, 0xD6)
C_ORG_D  = RGBColor(0x84, 0x3C, 0x0B)
C_GRAY_L = RGBColor(0xF5, 0xF5, 0xF5)
C_GRAY_B = RGBColor(0xDD, 0xDD, 0xDD)

OUT = "/mnt/d/code/study/research-protocol/毕设/开题PPT"
RECT = 1  # no rounding


def add_run(p, text, sz=18, bold=False, color=C_MED):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(sz)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = FONT


def main():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    box_l = Inches(0.5)
    box_w = Inches(12.3)

    # ── Title ──
    tb = slide.shapes.add_textbox(box_l, Inches(0.2), box_w, Inches(0.5))
    tb.text_frame.word_wrap = True
    p = tb.text_frame.paragraphs[0]
    add_run(p, "1.2  国内外研究现状", sz=22, bold=True, color=C_DARK)

    # ── Top: summary banner ──
    ban_y = Inches(0.85)
    ban_h = Inches(0.7)
    ban = slide.shapes.add_shape(RECT, box_l, ban_y, box_w, ban_h)
    ban.fill.solid()
    ban.fill.fore_color.rgb = C_BLUE
    ban.line.fill.background()

    tb = slide.shapes.add_textbox(box_l + Inches(0.2), ban_y + Inches(0.1),
                                   box_w - Inches(0.4), ban_h - Inches(0.2))
    tb.text_frame.word_wrap = True
    p = tb.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    add_run(p, "信道估计与载波同步在 FSO 湍流场景中研究仍不完善", sz=20, bold=True, color=C_WHITE)

    # ── Main container box ──
    main_y = Inches(1.8)
    main_h = Inches(4.4)
    main_box = slide.shapes.add_shape(RECT, box_l, main_y, box_w, main_h)
    main_box.fill.solid()
    main_box.fill.fore_color.rgb = C_WHITE
    main_box.line.color.rgb = C_GRAY_B
    main_box.line.width = Pt(1.5)

    # ── Center vertical divider ──
    center_x = box_l + Inches(6.15)
    div = slide.shapes.add_shape(RECT, center_x, main_y + Inches(0.05),
                                  Inches(0.03), main_h - Inches(0.1))
    div.fill.solid()
    div.fill.fore_color.rgb = C_GRAY_B
    div.line.fill.background()

    # ── Left column: blue ──
    col_w = Inches(5.7)
    col_l = box_l + Inches(0.15)
    col_r = center_x + Inches(0.2)

    # Left header bar (solid dark)
    lh_y = main_y + Inches(0.12)
    lh_h = Inches(0.6)
    lh = slide.shapes.add_shape(RECT, col_l, lh_y, col_w, lh_h)
    lh.fill.solid()
    lh.fill.fore_color.rgb = C_BLUE_D
    lh.line.fill.background()

    tb = slide.shapes.add_textbox(col_l + Inches(0.15), lh_y + Inches(0.1),
                                   col_w - Inches(0.3), lh_h - Inches(0.2))
    tb.text_frame.word_wrap = True
    p = tb.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    add_run(p, "CE→CPR 级联影响未被定量分析", sz=18, bold=True, color=C_WHITE)

    # Left bullet cards
    left_bullets = [
        "CE 精度与下游 CPR 性能缺乏定量关联",
        "不同同步方法对估计误差灵敏度差异大",
        "设计者无法确定估计精度做到多少才能满足同步需求",
    ]
    bc_w = col_w - Inches(0.1)
    bc_h = Inches(0.95)
    bc_gap = Inches(0.15)
    bc_x = col_l + Inches(0.05)
    bc_start_y = lh_y + lh_h + Inches(0.18)
    for i, b in enumerate(left_bullets):
        by = bc_start_y + i * (bc_h + bc_gap)
        # Card background
        c = slide.shapes.add_shape(RECT, bc_x, by, bc_w, bc_h)
        c.fill.solid()
        c.fill.fore_color.rgb = C_BLUE_L
        c.line.color.rgb = C_BLUE_B
        c.line.width = Pt(1)
        # Left accent
        s = slide.shapes.add_shape(RECT, bc_x, by, Inches(0.06), bc_h)
        s.fill.solid()
        s.fill.fore_color.rgb = C_BLUE
        s.line.fill.background()
        # Text
        tb = slide.shapes.add_textbox(bc_x + Inches(0.2), by + Inches(0.15),
                                       bc_w - Inches(0.35), bc_h - Inches(0.3))
        tb.text_frame.word_wrap = True
        p = tb.text_frame.paragraphs[0]
        add_run(p, b, sz=18, color=C_DARK)

    # ── Right column: orange ──
    rh = slide.shapes.add_shape(RECT, col_r, lh_y, col_w, lh_h)
    rh.fill.solid()
    rh.fill.fore_color.rgb = C_ORG_D
    rh.line.fill.background()

    tb = slide.shapes.add_textbox(col_r + Inches(0.15), lh_y + Inches(0.1),
                                   col_w - Inches(0.3), lh_h - Inches(0.2))
    tb.text_frame.word_wrap = True
    p = tb.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    add_run(p, "载波同步湍流适配设计方法缺失", sz=18, bold=True, color=C_WHITE)

    right_bullets = [
        "经典设计公式多基于白噪声，湍流下设计方法不完善",
        "性能分析仅涉及少数方法，未将 CE 误差纳入",
        "缺乏分湍流条件的方法选择与参数设计方法",
    ]
    rc_x = col_r + Inches(0.05)
    for i, b in enumerate(right_bullets):
        by = bc_start_y + i * (bc_h + bc_gap)
        c = slide.shapes.add_shape(RECT, rc_x, by, bc_w, bc_h)
        c.fill.solid()
        c.fill.fore_color.rgb = C_ORG_L
        c.line.color.rgb = C_ORG_B
        c.line.width = Pt(1)
        s = slide.shapes.add_shape(RECT, rc_x, by, Inches(0.06), bc_h)
        s.fill.solid()
        s.fill.fore_color.rgb = C_ORG
        s.line.fill.background()
        tb = slide.shapes.add_textbox(rc_x + Inches(0.2), by + Inches(0.15),
                                       bc_w - Inches(0.35), bc_h - Inches(0.3))
        tb.text_frame.word_wrap = True
        p = tb.text_frame.paragraphs[0]
        add_run(p, b, sz=18, color=C_DARK)

    # ── Bottom: bridge bar ──
    bridge_y = main_y + main_h + Inches(0.2)
    bridge_h = Inches(0.65)
    bridge = slide.shapes.add_shape(RECT, box_l, bridge_y, box_w, bridge_h)
    bridge.fill.solid()
    bridge.fill.fore_color.rgb = C_GRAY_L
    bridge.line.color.rgb = C_BLUE_B
    bridge.line.width = Pt(1.5)

    # Left accent on bridge
    bs = slide.shapes.add_shape(RECT, box_l, bridge_y, Inches(0.08), bridge_h)
    bs.fill.solid()
    bs.fill.fore_color.rgb = C_BLUE
    bs.line.fill.background()

    # Right accent on bridge
    bs2 = slide.shapes.add_shape(RECT, box_l + box_w - Inches(0.08), bridge_y,
                                  Inches(0.08), bridge_h)
    bs2.fill.solid()
    bs2.fill.fore_color.rgb = C_ORG
    bs2.line.fill.background()

    tb = slide.shapes.add_textbox(box_l + Inches(0.3), bridge_y + Inches(0.1),
                                   box_w - Inches(0.6), bridge_h - Inches(0.2))
    tb.text_frame.word_wrap = True
    p = tb.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    add_run(p, "估计误差级联 ", sz=18, bold=True, color=C_BLUE)
    add_run(p, "→ ", sz=20, bold=True, color=C_DARK)
    add_run(p, " 约束同步设计", sz=18, bold=True, color=C_ORG)

    out = os.path.join(OUT, "P07-研究空白分析v4.pptx")
    prs.save(out)
    print(f"Saved → {out}")


if __name__ == "__main__":
    main()
