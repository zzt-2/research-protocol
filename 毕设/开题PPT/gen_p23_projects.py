#!/usr/bin/env python3
"""P23: 前期相关项目经验 — 上方项目卡片+特色 + 下方2×2共通技术卡片（直角）"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import os

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
FONT = "微软雅黑"

C_DARK   = RGBColor(0x33, 0x33, 0x33)
C_MED    = RGBColor(0x55, 0x55, 0x55)
C_BLUE   = RGBColor(0x2E, 0x75, 0xB6)
C_BLUE_L = RGBColor(0xDE, 0xEB, 0xF7)
C_BLUE_B = RGBColor(0x9D, 0xC3, 0xE6)
C_ORG    = RGBColor(0xC5, 0x5A, 0x11)
C_ORG_L  = RGBColor(0xFF, 0xF8, 0xE1)
C_ORG_B  = RGBColor(0xFB, 0xE5, 0xD6)
C_GRAY_L = RGBColor(0xF8, 0xF8, 0xF8)
C_GRAY_B = RGBColor(0xDD, 0xDD, 0xDD)
C_WHITE  = RGBColor(0xFF, 0xFF, 0xFF)

OUT = "/mnt/d/code/study/research-protocol/毕设/开题PPT"

# shape type 1 = rectangle (no rounding)
RECT = 1


def add_run(p, text, sz=18, bold=False, color=C_MED):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(sz)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = FONT


def add_title_card(slide, left, top, w, h, number, title, subtitle, accent_c):
    card = slide.shapes.add_shape(RECT, left, top, w, h)
    card.fill.solid()
    card.fill.fore_color.rgb = C_WHITE
    card.line.color.rgb = C_GRAY_B
    card.line.width = Pt(1)

    bar = slide.shapes.add_shape(RECT, left, top, w, Inches(0.05))
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent_c
    bar.line.fill.background()

    tb = slide.shapes.add_textbox(left + Inches(0.15), top + Inches(0.1),
                                   w - Inches(0.3), h - Inches(0.15))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    add_run(p, f"{number} ", sz=18, bold=True, color=accent_c)
    add_run(p, title, sz=18, bold=True, color=C_DARK)
    p2 = tf.add_paragraph()
    add_run(p2, subtitle, sz=13, color=C_MED)


def add_feature_box(slide, left, top, w, h, accent_c, text):
    box = slide.shapes.add_shape(RECT, left, top, w, h)
    box.fill.solid()
    box.fill.fore_color.rgb = C_GRAY_L
    box.line.color.rgb = accent_c
    box.line.width = Pt(1)

    strip = slide.shapes.add_shape(RECT, left, top, Inches(0.08), h)
    strip.fill.solid()
    strip.fill.fore_color.rgb = accent_c
    strip.line.fill.background()

    tb = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.08),
                                   w - Inches(0.3), h - Inches(0.16))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    add_run(p, text, sz=18, color=C_DARK)


def add_tech_card(slide, left, top, w, h, title, desc, accent_c, fill_c):
    card = slide.shapes.add_shape(RECT, left, top, w, h)
    card.fill.solid()
    card.fill.fore_color.rgb = fill_c
    card.line.color.rgb = accent_c
    card.line.width = Pt(1.5)

    bar = slide.shapes.add_shape(RECT, left, top, w, Inches(0.04))
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent_c
    bar.line.fill.background()

    tb = slide.shapes.add_textbox(left + Inches(0.15), top + Inches(0.12),
                                   w - Inches(0.3), h - Inches(0.2))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    add_run(p, title, sz=18, bold=True, color=accent_c)
    p.space_after = Pt(4)
    p2 = tf.add_paragraph()
    add_run(p2, desc, sz=18, color=C_MED)


def main():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # ── Title ──
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.5))
    tb.text_frame.word_wrap = True
    p = tb.text_frame.paragraphs[0]
    add_run(p, "前期相关项目经验", sz=22, bold=True, color=C_DARK)

    # ── Row 1: project title cards ──
    cw = Inches(3.9)
    gap = Inches(0.25)
    x1 = Inches(0.5)
    x2 = x1 + cw + gap
    x3 = x2 + cw + gap
    card_y = Inches(1.0)
    card_h = Inches(0.7)

    add_title_card(slide, x1, card_y, cw, card_h,
        "①", "卫星基带收发系统", "BPSK / QPSK · 多速率", C_BLUE)
    add_title_card(slide, x2, card_y, cw, card_h,
        "②", "星间链路收发信机", "BPSK / QPSK · 10Gsps", C_ORG)
    add_title_card(slide, x3, card_y, cw, card_h,
        "③", "卫星高速调制解调器", "QPSK → 64QAM", C_BLUE)

    # ── Feature boxes (under each card) ──
    feat_y = card_y + card_h + Inches(0.15)
    feat_h = Inches(1.5)

    add_feature_box(slide, x1, feat_y, cw, feat_h, C_BLUE,
        "两种xxx：RS / LDPC")
    add_feature_box(slide, x2, feat_y, cw, feat_h, C_ORG,
        "双向精密测距")
    add_feature_box(slide, x3, feat_y, cw, feat_h, C_BLUE,
        "多调制格式支持（QPSK → 64QAM），优化 Gardner 定时，软解调")

    # ── 共通技术基础 section label ──
    sec_y = feat_y + feat_h + Inches(0.3)
    sec = slide.shapes.add_textbox(Inches(0.5), sec_y, Inches(12.3), Inches(0.4))
    sec.text_frame.word_wrap = True
    p = sec.text_frame.paragraphs[0]
    add_run(p, "共通技术基础", sz=18, bold=True, color=C_BLUE)

    # ── 2×2 tech cards ──
    tc_w = Inches(5.95)
    tc_h = Inches(1.25)
    tc_gap_x = Inches(0.25)
    tc_gap_y = Inches(0.2)
    tc_x1 = Inches(0.5)
    tc_x2 = tc_x1 + tc_w + tc_gap_x
    tc_y1 = sec_y + Inches(0.45)
    tc_y2 = tc_y1 + tc_h + tc_gap_y

    add_tech_card(slide, tc_x1, tc_y1, tc_w, tc_h,
        "载波同步", "FFT 粗补偿 + DPLL / VV", C_BLUE, C_BLUE_L)
    add_tech_card(slide, tc_x2, tc_y1, tc_w, tc_h,
        "定时同步", "Gardner + 三次插值", C_ORG, C_ORG_L)
    add_tech_card(slide, tc_x1, tc_y2, tc_w, tc_h,
        "均衡", "CMA 盲均衡 / LMS 自适应均衡", C_ORG, C_ORG_L)
    add_tech_card(slide, tc_x2, tc_y2, tc_w, tc_h,
        "FPGA 全流程", "调制 → 解调 → 同步 → 均衡 → 编译码", C_BLUE, C_BLUE_L)

    out = os.path.join(OUT, "P23-前期项目经验v5.pptx")
    prs.save(out)
    print(f"Saved → {out}")


if __name__ == "__main__":
    main()
