#!/usr/bin/env python3
"""P22: 研究基础 - 两个小截图 + 已完成工作汇总"""

from pptx import Presentation
from pptx.util import Inches, Pt, Cm
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
import os

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
FONT = "微软雅黑"

C_DARK = RGBColor(0x44, 0x44, 0x44)
C_MED  = RGBColor(0x66, 0x66, 0x66)
C_BLUE = RGBColor(0x2E, 0x75, 0xB6)
C_BLUE_L = RGBColor(0xDE, 0xEB, 0xF7)
C_BLUE_B = RGBColor(0x9D, 0xC3, 0xE6)
C_ORG  = RGBColor(0xC5, 0x5A, 0x11)
C_ORG_L = RGBColor(0xFF, 0xF8, 0xE1)
C_ORG_B = RGBColor(0xFB, 0xE5, 0xD6)
C_GRAY_L = RGBColor(0xF5, 0xF5, 0xF5)
C_GRAY_B = RGBColor(0xDD, 0xDD, 0xDD)
C_PH_BG = RGBColor(0xF0, 0xF0, 0xF0)
C_PH_TXT = RGBColor(0xAA, 0xAA, 0xAA)
C_GRAY_LINE = RGBColor(0xCC, 0xCC, 0xCC)

OUT = "/mnt/d/code/study/research-protocol/毕设/开题PPT"


def add_run(p, text, sz=18, bold=False, color=C_MED):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(sz)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = FONT


def set_dash(shape):
    ln = shape.line._ln
    for existing in ln.findall(qn('a:prstDash')):
        ln.remove(existing)
    ln.append(ln.makeelement(qn('a:prstDash'), {'val': 'dash'}))


def add_placeholder(slide, left, top, w, h, text):
    shape = slide.shapes.add_shape(1, left, top, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = C_PH_BG
    shape.line.color.rgb = C_GRAY_LINE
    shape.line.width = Pt(1.5)
    set_dash(shape)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_top = Inches(0.4)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    add_run(p, text, sz=18, color=C_PH_TXT)


def add_card(slide, left, top, w, h, accent_c, number, title, desc):
    # Card background
    card = slide.shapes.add_shape(5, left, top, w, h)
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(0xFA, 0xFA, 0xFA)
    card.line.color.rgb = C_GRAY_B
    card.line.width = Pt(1)

    # Left accent strip
    strip = slide.shapes.add_shape(1, left, top, Inches(0.1), h)
    strip.fill.solid()
    strip.fill.fore_color.rgb = accent_c
    strip.line.fill.background()

    # Text
    tb = slide.shapes.add_textbox(
        left + Inches(0.25), top + Inches(0.1),
        w - Inches(0.35), h - Inches(0.2))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    add_run(p, f"{number} ", sz=18, bold=True, color=accent_c)
    add_run(p, title, sz=18, bold=True, color=C_DARK)
    p.space_after = Pt(4)
    p2 = tf.add_paragraph()
    add_run(p2, desc, sz=18, color=C_MED)


def main():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Title
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.7))
    tb.text_frame.word_wrap = True
    p = tb.text_frame.paragraphs[0]
    add_run(p, "研究基础与前期工作", sz=22, bold=True, color=C_DARK)

    # ── Row 1: two small screenshot cards side by side ──
    # Each card: [placeholder ~5cm] + [label + bullets]
    ph_w = Cm(5)   # ~1.97"
    ph_h = Cm(4)   # ~1.57"

    # Left card: MATLAB
    card1_x = Inches(0.5)
    card1_y = Inches(1.2)
    # screenshot
    add_placeholder(slide, card1_x, card1_y, ph_w, ph_h, "MATLAB\n仿真截图")
    # label + bullets (right of screenshot)
    tb1 = slide.shapes.add_textbox(
        card1_x + ph_w + Inches(0.2), card1_y,
        Inches(5.5) - ph_w - Inches(0.2), ph_h)
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    add_run(p, "仿真平台（MATLAB）", sz=18, bold=True, color=C_BLUE)
    p.space_after = Pt(4)
    for txt in [
        "QPSK 相干检测链路级仿真平台",
        "含 GG 湍流信道、导频辅助 CE、VV/BPS/DPLL 同步",
        "已完成 Monte Carlo 参数扫描",
    ]:
        pb = tf1.add_paragraph()
        add_run(pb, f"• {txt}", sz=18, color=C_MED)
        pb.space_after = Pt(2)

    # Right card: Vivado
    card2_x = Inches(6.85)
    card2_y = Inches(1.2)
    add_placeholder(slide, card2_x, card2_y, ph_w, ph_h, "Vivado\n仿真截图")
    tb2 = slide.shapes.add_textbox(
        card2_x + ph_w + Inches(0.2), card2_y,
        Inches(5.5) - ph_w - Inches(0.2), ph_h)
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    add_run(p, "硬件平台（FPGA）", sz=18, bold=True, color=C_ORG)
    p.space_after = Pt(4)
    for txt in [
        "载波同步链 RTL 设计（FFT-FOE + DPLL）",
        "通过功能仿真验证",
        "目标平台：Xilinx Zynq 系列",
    ]:
        pb = tf2.add_paragraph()
        add_run(pb, f"• {txt}", sz=18, color=C_MED)
        pb.space_after = Pt(2)

    # ── Row 2: 2×2 cards for completed work ──
    # Section title
    sec = slide.shapes.add_textbox(Inches(0.5), Inches(3.5), Inches(12.3), Inches(0.45))
    sec.text_frame.word_wrap = True
    p = sec.text_frame.paragraphs[0]
    add_run(p, "已完成的初步工作", sz=18, bold=True, color=C_DARK)

    cw = Inches(6.0)
    ch = Inches(1.35)
    x1, x2 = Inches(0.5), Inches(6.8)
    y1, y2 = Inches(4.05), Inches(5.55)

    add_card(slide, x1, y1, cw, ch, C_BLUE, "①", "载波同步对比",
             "三种湍流条件下 VV/BPS/DPLL 的 BER 性能对比仿真")
    add_card(slide, x2, y1, cw, ch, C_BLUE, "②", "估计精度分析",
             "NMSE 随 SNR 与导频密度变化的规律分析")
    add_card(slide, x1, y2, cw, ch, C_ORG, "③", "DPLL 参数设计",
             "环路带宽与阻尼系数对跟踪性能的影响分析")
    add_card(slide, x2, y2, cw, ch, C_ORG, "④", "调制格式差异",
             "QPSK 与 16-QAM 对估计误差的灵敏度差异分析")

    out = os.path.join(OUT, "P22-研究基础v3.pptx")
    prs.save(out)
    print(f"Saved → {out}")


if __name__ == "__main__":
    main()
