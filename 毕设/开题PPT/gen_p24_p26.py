#!/usr/bin/env python3
"""P24-P26: 三张仿真结果页（嵌入 Ch4 技术路线）

统一模板：标题 + 图（~55% 页高）+ 结论句
橙色主题对齐 Ch4

Usage: cd /mnt/d/code/study/research-protocol/毕设/开题PPT && ~/.venvs/torch/bin/python gen_p24_p26.py
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
import os

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
FONT = "微软雅黑"

C_DARK  = RGBColor(0x33, 0x33, 0x33)
C_ORG   = RGBColor(0xC5, 0x5A, 0x11)
C_ORG_D = RGBColor(0x84, 0x3C, 0x0B)
C_ORG_L = RGBColor(0xFF, 0xF8, 0xE1)

FIG_DIR = "/mnt/d/code/study/research-protocol/毕设/开题报告/figures/png"
OUT_DIR = "/mnt/d/code/study/research-protocol/毕设/开题PPT"


def add_run(p, text, sz=18, bold=False, color=C_DARK):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(sz)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = FONT


def add_title(slide, text):
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.15), Inches(12.3), Inches(0.65))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    add_run(p, text, sz=20, bold=True, color=C_DARK)


def add_figure(slide, img_path):
    slide.shapes.add_picture(
        img_path,
        left=Inches(0.5),
        top=Inches(0.95),
        width=Inches(12.3),
    )


def add_conclusion(slide, text):
    shape = slide.shapes.add_shape(
        1,  # rectangle
        Inches(0.5), Inches(5.35),
        Inches(12.3), Inches(0.85),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = C_ORG_L
    shape.line.fill.background()

    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.2)
    tf.margin_right = Inches(0.2)
    tf.margin_top = Inches(0.1)
    tf.margin_bottom = Inches(0.1)

    p = tf.paragraphs[0]
    add_run(p, "→  " + text, sz=18, color=C_ORG_D)


def add_label(slide, text):
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(6.35), Inches(12.3), Inches(0.4))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    add_run(p, text, sz=14, color=RGBColor(0x99, 0x99, 0x99))


SLIDES = [
    {
        "title": "典型参数设置下，三种载波同步方法在强湍流时性能明显分化",
        "fig": "ber-snr-carrier-sync.png",
        "conclusion": "典型参数下湍流越强方法间差距越大；强湍流时 VV/BPS 性能趋近随机水平，DPLL 仍保持较低误码率",
        "label": "Ch4 初步验证 ①/③  ·  数据：VV/BPS/DPLL BER vs SNR  ·  三湍流条件",
    },
    {
        "title": "前馈方法对平均窗长敏感，优化参数后可逼近 DPLL 性能",
        "fig": "ber-nw-sensitivity.png",
        "conclusion": "湍流越强所需窗长越大；优化参数后 VV 可逼近 DPLL，表明差异源于参数敏感度而非方法本身",
        "label": "Ch4 初步验证 ②/③  ·  数据：VV/BPS BER vs 平均窗长 Nw  ·  三湍流条件",
    },
    {
        "title": "DPLL 参数宽容度更高，但强湍流下 ωn 过大导致周期跳跃",
        "fig": "ber-dpll-omega-sensitivity.png",
        "conclusion": "弱/中等湍流下 DPLL 对 ωn 宽容度高；强湍流下 ωn 过大导致性能急剧恶化，不同湍流条件存在最优 ωn 区间",
        "label": "Ch4 初步验证 ③/③  ·  数据：DPLL BER vs 环路带宽 ωn  ·  三湍流 × 三 SNR",
    },
]


def main():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    for spec in SLIDES:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        add_title(s, spec["title"])
        add_figure(s, os.path.join(FIG_DIR, spec["fig"]))
        add_conclusion(s, spec["conclusion"])
        add_label(s, spec["label"])

    out = os.path.join(OUT_DIR, "P24-P26-仿真结果.pptx")
    prs.save(out)
    print(f"Saved {len(prs.slides)} slides → {out}")


if __name__ == "__main__":
    main()
